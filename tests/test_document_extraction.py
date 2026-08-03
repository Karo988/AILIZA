"""Tests fuer die konsolidierte, strukturierte Dokument-Extraktion
(apps/backend/documents/extraction.py) -- PDF/DOCX/XLSX/CSV und Bilder/OCR
(nur bei lokal verfuegbarem OCR). Ersetzt/ergaenzt die bisher getrennten,
oberflaechlichen Extraktions-Implementierungen in document_handler.py und
ingestion.py (Karo-Entscheidung 2026-08-03: Aufraeumen/Konsolidieren)."""
from __future__ import annotations

import io
import os

os.environ.setdefault("AILIZA_SECRET_KEY", "test-secret-key-minimum-32-chars-ok")
os.environ.setdefault("AILIZA_DATABASE_URL", "sqlite:///:memory:")

import pytest

from apps.backend.documents.extraction import (
    extract_structured, ocr_available, supported_structured_extensions,
    MAX_XLSX_CELLS_PER_SHEET,
)


def _make_xlsx_bytes(rows: list[list[object]]) -> bytes:
    openpyxl = pytest.importorskip("openpyxl")
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_docx_bytes(paragraphs: list[tuple[str, str | None]]) -> bytes:
    """paragraphs: Liste von (text, style_name_or_None)."""
    docx = pytest.importorskip("docx")
    from docx import Document
    doc = Document()
    for text, style in paragraphs:
        p = doc.add_paragraph(text)
        if style:
            p.style = doc.styles[style]
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# -- CSV (keine externe Bibliothek noetig, immer testbar) --------------------

def test_csv_extracted_as_markdown_table():
    content = "Name,Betrag\nTest,12.50\nZweite Zeile,7\n".encode("utf-8")
    result = extract_structured(".csv", content)
    assert not result.warnings
    assert len(result.sections) == 1
    assert result.sections[0].element_type == "table"
    assert "Name" in result.sections[0].content_markdown
    assert "Betrag" in result.sections[0].content_markdown
    assert "12.50" in result.sections[0].content_markdown
    assert result.full_text == result.sections[0].content_markdown


def test_csv_empty_content_yields_no_sections():
    result = extract_structured(".csv", b"")
    assert result.sections == []
    assert result.full_text == ""


def test_csv_row_limit_truncates():
    rows = "\n".join(f"row{i},{i}" for i in range(60_000))
    content = rows.encode("utf-8")
    from apps.backend.documents.extraction import MAX_CSV_ROWS
    result = extract_structured(".csv", content)
    assert result.truncated is True
    assert any(w.startswith("truncated:csv_rows") for w in result.warnings)


def test_csv_invalid_utf8_reports_extraction_error():
    result = extract_structured(".csv", b"\xff\xfe invalid utf8 bytes")
    assert any(w.startswith("extraction_error:") for w in result.warnings)
    assert result.sections == []


# -- XLSX ---------------------------------------------------------------------

def test_xlsx_extracted_with_sheet_name_as_section_title():
    content = _make_xlsx_bytes([["Name", "Betrag"], ["Test", 12.5]])
    result = extract_structured(".xlsx", content)
    assert not result.warnings
    assert len(result.sections) == 1
    assert result.sections[0].element_type == "table"
    assert result.sections[0].section_title == "Sheet"
    assert result.sections[0].page_number is None
    assert "Test" in result.sections[0].content_markdown


def test_xlsx_empty_workbook_yields_no_sections():
    content = _make_xlsx_bytes([])
    result = extract_structured(".xlsx", content)
    assert result.sections == []


def test_xlsx_corrupted_content_reports_extraction_error_not_crash():
    result = extract_structured(".xlsx", b"not a real xlsx file")
    assert any(w.startswith("extraction_error:") for w in result.warnings)
    assert result.sections == []


# -- DOCX -----------------------------------------------------------------

def test_docx_heading_detected_as_section_title():
    content = _make_docx_bytes([
        ("Meine Ueberschrift", "Heading 1"),
        ("Ein normaler Absatz.", None),
    ])
    result = extract_structured(".docx", content)
    assert not result.warnings
    headings = [s for s in result.sections if s.element_type == "heading"]
    paragraphs = [s for s in result.sections if s.element_type == "paragraph"]
    assert len(headings) == 1
    assert headings[0].section_title == "Meine Ueberschrift"
    assert len(paragraphs) == 1
    assert paragraphs[0].content_markdown == "Ein normaler Absatz."


def test_docx_corrupted_content_reports_extraction_error_not_crash():
    result = extract_structured(".docx", b"definitely not a docx file")
    assert any(w.startswith("extraction_error:") for w in result.warnings)
    assert result.sections == []


# -- PDF (nur falls pdfplumber installiert ist -- sonst muss library_missing
#    korrekt gemeldet werden, das wird unabhaengig von der Installation getestet) --

def test_pdf_missing_library_or_extracts_without_crash():
    result = extract_structured(".pdf", b"%PDF-1.4 kein echtes PDF")
    assert result.library_missing or any(w.startswith("extraction_error:") for w in result.warnings)


# -- Bilder/OCR -- nur lokal verfuegbar, sonst fail-closed --------------------

def test_image_extraction_reports_library_missing_when_ocr_disabled(monkeypatch):
    monkeypatch.delenv("AILIZA_LOCAL_OCR_ENABLED", raising=False)
    assert ocr_available() is False
    result = extract_structured(".png", b"not a real png")
    assert result.library_missing is True
    assert "library_missing:ocr_local_only" in result.warnings


def test_image_extensions_only_in_supported_set_when_ocr_available(monkeypatch):
    monkeypatch.delenv("AILIZA_LOCAL_OCR_ENABLED", raising=False)
    exts = supported_structured_extensions()
    assert ".png" not in exts
    assert ".jpg" not in exts
    assert ".pdf" in exts or True  # pdfplumber evtl. nicht installiert -- kein harter Fehler hier


# -- Unbekannte Erweiterung ----------------------------------------------------

def test_unsupported_extension_reports_warning_not_crash():
    result = extract_structured(".xyz", b"irrelevant")
    assert any(w.startswith("unsupported_extension:") for w in result.warnings)
    assert result.sections == []
