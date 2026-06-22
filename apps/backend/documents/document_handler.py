"""
AILIZA Dokumenten-Workflow MVP
=============================
Kein Upload ohne Governance. Whitelist + Groessenpruefung + Klassifikation +
Gate 6 Prompt-Injection-Erkennung.

Gate 6 Prinzip:
  Dokumentinhalt ist immer Nutzinhalt — niemals Steuerbefehl.
  Eingebettete Anweisungen werden erkannt, als SECURITY_SENSITIVE markiert und
  blockieren die automatische externe Verarbeitung. Policy-Gates bleiben immer
  hoeher priorisiert als jeder Dokumenttext.
"""
from __future__ import annotations

import os
import re
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone

try:
    from ..governance.data_governance import ClassificationResult, DataClass, classify
    from ..governance.data_matrix import PolicyDecision, check_data_target
    from ..governance.data_governance import DataTarget
except ImportError:  # pragma: no cover
    from governance.data_governance import ClassificationResult, DataClass, classify, DataTarget
    from governance.data_matrix import PolicyDecision, check_data_target


ALLOWED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".txt", ".csv"}
MAX_FILE_SIZE_MB = 10
_RETENTION_DAYS = int(os.getenv("AILIZA_DOCUMENT_RETENTION_DAYS", "30"))

# Gate 6 — Prompt-Injection-Muster (case-insensitive, Pattern-basiert, kein LLM)
_INJECTION_PATTERNS: list[re.Pattern[str]] = [
    re.compile(r"ignore\s+(?:all\s+)?(?:previous\s+)?instructions?", re.I),
    re.compile(r"ignoriere\s+(?:alle\s+)?(?:vorherigen?\s+)?anweisungen?", re.I),
    re.compile(r"du\s+bist\s+jetzt", re.I),
    re.compile(r"you\s+are\s+now", re.I),
    re.compile(r"(?:^|\s)system\s*:", re.I | re.M),
    re.compile(r"(?:^|\s)developer\s*:", re.I | re.M),
    re.compile(r"\[INST\]", re.I),
    re.compile(r"\[SYS\]", re.I),
    re.compile(r"<\|system\|>", re.I),
    re.compile(r"override\s+policy", re.I),
    re.compile(r"bypass\s+governance", re.I),
    re.compile(r"ignore\s+rules?", re.I),
    re.compile(r"act\s+as\s+(?:dan|jailbreak|unrestricted|admin|root)", re.I),
    re.compile(r"you\s+(?:must|shall|will)\s+(?:now\s+)?(?:ignore|forget|disregard)", re.I),
]


@dataclass
class InjectionScanResult:
    """Ergebnis des Gate-6-Scans — enthält nie den Originaltext."""
    injection_detected: bool
    pattern_count: int  # Anzahl Treffer, kein Treffertext


def _scan_for_injection(text: str) -> InjectionScanResult:
    """
    Gate 6: Prüft ob Text Prompt-Injection-Muster enthält.
    Gibt nur technische Metadaten zurück — niemals den erkannten Text.
    """
    if not text:
        return InjectionScanResult(injection_detected=False, pattern_count=0)
    count = sum(1 for p in _INJECTION_PATTERNS if p.search(text))
    return InjectionScanResult(injection_detected=count > 0, pattern_count=count)


@dataclass
class DocumentScanResult:
    allowed: bool
    file_type: str
    size_bytes: int
    classification: ClassificationResult
    decision: str  # allow / redact / approval_required / block
    reason: str
    expires_at: str
    injection_detected: bool = False      # Gate 6
    injection_pattern_count: int = 0      # Gate 6 — kein Treffertext

    def to_audit_dict(self) -> dict:
        """Audit-Light: keine Rohinhalte, keine Treffer-Texte."""
        return {
            "file_type": self.file_type,
            "size_bytes": self.size_bytes,
            "decision": self.decision,
            "highest_risk_class": self.classification.highest_risk_class,
            "needs_review": self.classification.needs_review,
            "injection_detected": self.injection_detected,
            "injection_pattern_count": self.injection_pattern_count,
        }


def _extract_text(ext: str, content: bytes) -> str:
    if ext in {".txt", ".csv"}:
        try:
            return content.decode("utf-8", errors="ignore")
        except Exception:
            return ""

    if ext == ".pdf":
        try:
            import io
            import pdfplumber
            with pdfplumber.open(io.BytesIO(content)) as pdf:
                return "\n".join(page.extract_text() or "" for page in pdf.pages)
        except ImportError:
            pass  # pdfplumber nicht installiert — Fallback
        except Exception:
            return ""

    if ext == ".docx":
        try:
            import io
            from docx import Document
            doc = Document(io.BytesIO(content))
            return "\n".join(p.text for p in doc.paragraphs)
        except ImportError:
            pass  # python-docx nicht installiert — Fallback
        except Exception:
            return ""

    if ext == ".xlsx":
        try:
            import io
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            parts: list[str] = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    parts.append(" ".join(str(c) for c in row if c is not None))
            return "\n".join(parts)
        except ImportError:
            pass  # openpyxl nicht installiert — Fallback
        except Exception:
            return ""

    return ""


def scan_document(filename: str, content: bytes) -> DocumentScanResult:
    ext = os.path.splitext(filename or "")[1].lower()
    size = len(content or b"")
    expires_at = (datetime.now(timezone.utc) + timedelta(days=_RETENTION_DAYS)).isoformat()

    # 1. Whitelist
    if ext not in ALLOWED_EXTENSIONS:
        return DocumentScanResult(
            allowed=False, file_type=ext or "unknown", size_bytes=size,
            classification=classify(""), decision="block",
            reason=f"Dateityp {ext or 'unbekannt'} ist nicht erlaubt.", expires_at=expires_at)

    # 2. Groesse
    if size > MAX_FILE_SIZE_MB * 1024 * 1024:
        return DocumentScanResult(
            allowed=False, file_type=ext, size_bytes=size,
            classification=classify(""), decision="block",
            reason=f"Datei groesser als {MAX_FILE_SIZE_MB} MB.", expires_at=expires_at)

    # 3. Text extrahieren
    text = _extract_text(ext, content)

    # 4. Gate 6 — Prompt-Injection-Erkennung (vor Klassifikation)
    injection = _scan_for_injection(text)
    if injection.injection_detected:
        # Dokumentinhalt mit Injection-Muster → SECURITY_SENSITIVE, needs_review, kein externer Call
        from dataclasses import replace as _replace
        inj_classification = ClassificationResult(
            data_classes=[DataClass.SECURITY_SENSITIVE],
            highest_risk_class=DataClass.SECURITY_SENSITIVE,
            needs_review=True,
            requires_human_decision=False,
        )
        return DocumentScanResult(
            allowed=False,
            file_type=ext,
            size_bytes=size,
            classification=inj_classification,
            decision="block",
            reason=(
                "Dokument enthält eingebettete Anweisungsmuster (Prompt-Injection-Verdacht). "
                "Externe Verarbeitung blockiert. Manuelle Prüfung erforderlich."
            ),
            expires_at=expires_at,
            injection_detected=True,
            injection_pattern_count=injection.pattern_count,
        )

    # 5. Klassifizieren
    classification = classify(text)

    # 6. Entscheidung (Ziel: externes LLM, ohne aktives ProviderProfile by default)
    decision = check_data_target(
        data_classes=classification.data_classes,
        target=DataTarget.EXTERNAL_LLM,
        redaction_applied=False,
        approval_given=False,
        provider_profile_active=False,
    )
    decision_map = {
        PolicyDecision.ALLOW: ("allow", "Dokument freigegeben."),
        PolicyDecision.ALLOW_WITH_NOTICE: ("allow", "Dokument freigegeben (mit Hinweis)."),
        PolicyDecision.REDACT_REQUIRED: ("redact", "Dokument muss anonymisiert werden."),
        PolicyDecision.APPROVAL_REQUIRED: ("approval_required", "Freigabe erforderlich."),
        PolicyDecision.BLOCK: ("block", "Dokument enthaelt nicht erlaubte Datenklassen."),
    }
    decision_str, reason = decision_map.get(decision, ("block", "Unklar — blockiert."))

    return DocumentScanResult(
        allowed=decision_str in {"allow", "redact", "approval_required"},
        file_type=ext, size_bytes=size, classification=classification,
        decision=decision_str, reason=reason, expires_at=expires_at,
        injection_detected=False, injection_pattern_count=0,
    )
